import json
import logging
import os
import re
import time
import httpx
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
import urllib.parse
from PIL import Image

from dotenv import load_dotenv
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from googleapiclient.http import MediaFileUpload
import os.path
from google_auth_oauthlib.flow import InstalledAppFlow
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from notion_client import Client
from io import BytesIO
from docx import Document
from docx.shared import Inches

# Load environment variables from .env file
load_dotenv()

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

BASE_DIR = Path(__file__).resolve().parent

NOTION_TOKEN = os.getenv("NOTION_TOKEN", "").strip()
DATABASE_ID = os.getenv("DATABASE_ID", "").strip()
DRIVE_FOLDER_ID = os.getenv("DRIVE_FOLDER_ID", "").strip()
SERVICE_ACCOUNT_FILE = os.getenv(
    "GOOGLE_APPLICATION_CREDENTIALS",
    str(BASE_DIR / "notion-sync-483309-02a27dfe1d63.json"),
)

USE_STATUS_FILTER = os.getenv("USE_STATUS_FILTER", "false").lower() in ("1", "true", "yes")
STATUS_PROPERTY = os.getenv("NOTION_STATUS_PROPERTY", "Status").strip()
STATUS_VALUE = os.getenv("NOTION_STATUS_VALUE", "Done").strip()
NOTION_TOKEN_V2 = os.getenv("NOTION_TOKEN_V2", "").strip()

OVERWRITE_EXISTING = os.getenv("OVERWRITE_EXISTING", "false").lower() in ("1", "true", "yes")

SYNC_STATE_PATH = BASE_DIR / "sync_state.json"
EXPORTS_DIR = BASE_DIR / "exports"

NOTION_RATE_LIMIT_DELAY = 0.4
DEFAULT_START_TIME = "2020-01-01T00:00:00.000Z"

DRIVE_SCOPES = ["https://www.googleapis.com/auth/drive"]


def load_last_sync_time() -> str:
    if SYNC_STATE_PATH.exists():
        try:
            data = json.loads(SYNC_STATE_PATH.read_text(encoding="utf-8"))
            last_sync_time = data.get("last_sync_time")
            if last_sync_time:
                return last_sync_time
        except (json.JSONDecodeError, OSError):
            pass
    return DEFAULT_START_TIME


def save_last_sync_time(iso_time: str) -> None:
    SYNC_STATE_PATH.write_text(
        json.dumps({"last_sync_time": iso_time}, indent=2),
        encoding="utf-8",
    )


def now_utc_iso() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def sanitize_filename(name: str, max_len: int = 120) -> str:
    cleaned = re.sub(r"[<>:\"/\\\\|?*]", "_", name)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    if not cleaned:
        cleaned = "Untitled"
    return cleaned[:max_len]


def extract_rich_text(rich_text) -> str:
    if not rich_text:
        return ""
    return "".join(part.get("plain_text", "") for part in rich_text)


def build_notion_filter(last_sync_time: str) -> dict:
    filters = []

    if USE_STATUS_FILTER and STATUS_PROPERTY and STATUS_VALUE:
        filters.append(
            {
                "property": STATUS_PROPERTY,
                "status": {"equals": STATUS_VALUE},
            }
        )

    filters.append(
        {
            "timestamp": "last_edited_time",
            "last_edited_time": {"after": last_sync_time},
        }
    )

    if len(filters) == 1:
        return filters[0]

    return {"or": filters}


def query_all_pages(notion: Client, database_id: str, last_sync_time: str) -> List[Dict[str, Any]]:
    """
    Use underlying request method to fetch all pages, bypassing version compatibility issues.
    """
    if len(database_id) == 32:
        database_id = (
            f"{database_id[:8]}-{database_id[8:12]}-{database_id[12:16]}-"
            f"{database_id[16:20]}-{database_id[20:]}"
        )
        logging.info(f"Formatted Database ID to UUID: {database_id}")

    all_pages = []
    has_more = True
    next_cursor = None

    # 1. Build filter (Payload)
    # Use existing helper to respect .env settings and last_sync_time
    notion_filter = build_notion_filter(last_sync_time)

    logging.info(f"Querying database {database_id}...")
    
    headers = {
        "Authorization": f"Bearer {NOTION_TOKEN}",
        "Notion-Version": "2022-06-28",
        "Content-Type": "application/json",
    }
    url = f"https://api.notion.com/v1/databases/{database_id}/query"

    # 2. Loop through pages
    while has_more:
        body = {
            "page_size": 100  # Fetch 100 items per request
        }

        # Only add filter if it exists
        if notion_filter:
            body["filter"] = notion_filter
        if next_cursor:
            body["start_cursor"] = next_cursor

        try:
            print(f"DEBUG BODY: {json.dumps(body, indent=2)}")
            sys.stdout.flush()
            # === FIX: Use httpx directly because notion-client v2.7.0 is broken ===
            response = httpx.post(url, headers=headers, json=body, timeout=60.0)
            response.raise_for_status()
            data = response.json()
            # ==========================================================

        except Exception as e:
            logging.error(f"Error querying Notion API: {e}")
            raise e

        # 3. Process results
        results = data.get("results", [])
        all_pages.extend(results)

        # Check for next page
        has_more = data.get("has_more", False)
        next_cursor = data.get("next_cursor")

        # Pause to avoid rate limits
        time.sleep(0.3)

    logging.info(f"Found {len(all_pages)} pages to sync.")
    return all_pages


def list_block_children(notion: Client, block_id: str) -> List[Dict[str, Any]]:
    blocks = []
    cursor = None

    while True:
        response = notion.blocks.children.list(
            block_id=block_id,
            start_cursor=cursor,
            page_size=100,
        )
        time.sleep(NOTION_RATE_LIMIT_DELAY)

        blocks.extend(response.get("results", []))
        if not response.get("has_more"):
            break
        cursor = response.get("next_cursor")

    return blocks


def download_image_to_stream(url: str, block_id: str = None) -> Optional[BytesIO]:
    """下载图片并返回 BytesIO (Robust Mode w/ Proxy Fallback)"""
    
    # === Common: Convert to PNG if needed for python-docx ===
    def convert_to_supported_format(data_bytes: bytes) -> BytesIO:
        try:
            img = Image.open(BytesIO(data_bytes))
            if img.format not in ['JPEG', 'PNG']:
                # Convert WebP etc to PNG
                out = BytesIO()
                img.convert("RGB").save(out, format="PNG")
                out.seek(0)
                return out
            else:
                return BytesIO(data_bytes)
        except Exception as e:
            logging.warning(f"Image conversion failed: {e}")
            return BytesIO(data_bytes) # Try original as last resort

    # === Direct Download ===
    try:
        # 使用 verify=False 和 headers 模拟浏览器，防止某些防盗链 (如小红书)
        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
        }
        # 小红书特殊处理
        if "xhscdn.com" in url:
             headers["Referer"] = "https://www.xiaohongshu.com/"

        # http2=False 更稳定
        with httpx.Client(http2=False, verify=False, headers=headers, follow_redirects=True, timeout=30.0) as client:
            response = client.get(url)
            if response.status_code == 200:
                return convert_to_supported_format(response.content)
            else:
                logging.warning(f"Direct download failed: {response.status_code}")
    except Exception as e:
        logging.warning(f"Failed to download image (direct): {e}")

    # === Fallback: Notion Proxy ===
    if block_id and NOTION_TOKEN_V2:
        logging.info(f"Attempting Fallback via Notion Proxy for block {block_id}...")
        try:
            encoded_url = urllib.parse.quote(url, safe='')
            proxy_url = f"https://www.notion.so/image/{encoded_url}?table=block&id={block_id}&cache=v2"
            
            proxy_headers = {
                "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
                "Cookie": f"token_v2={NOTION_TOKEN_V2}",
                "Accept": "image/webp,image/apng,image/*,*/*;q=0.8"
            }
            with httpx.Client(http2=False, verify=False, follow_redirects=True, timeout=30.0) as client:
                r = client.get(proxy_url, headers=proxy_headers)
                if r.status_code == 200:
                     logging.info(f"[OK] Proxy Download SUCCESS! Size: {len(r.content)}")
                     return convert_to_supported_format(r.content)
                else:
                    logging.warning(f"[FAIL] Proxy Download Failed: Status {r.status_code}")
        except Exception as e:
            logging.error(f"[FAIL] Proxy Download Exception: {e}")

    return None


def write_block_to_docx(doc: Document, block: Dict[str, Any], notion: Client, depth: int = 0) -> None:
    block_type = block.get("type")
    data = block.get(block_type, {})
    indent_level = depth  # Rudimentary indentation 

    # Helper to add indented text
    def add_text_paragraph(text, style=None):
        if not text:
            return
        p = doc.add_paragraph(text, style=style)
        if indent_level > 0 and style != 'List Bullet' and style != 'List Number':
             p.paragraph_format.left_indent = Inches(0.25 * indent_level)

    if block_type == "paragraph":
        text = extract_rich_text(data.get("rich_text"))
        add_text_paragraph(text)
    elif block_type == "heading_1":
        text = extract_rich_text(data.get("rich_text"))
        doc.add_heading(text, level=1)
    elif block_type == "heading_2":
        text = extract_rich_text(data.get("rich_text"))
        doc.add_heading(text, level=2)
    elif block_type == "heading_3":
        text = extract_rich_text(data.get("rich_text"))
        doc.add_heading(text, level=3)
    elif block_type == "bulleted_list_item":
        text = extract_rich_text(data.get("rich_text"))
        doc.add_paragraph(text, style='List Bullet')
    elif block_type == "numbered_list_item":
        text = extract_rich_text(data.get("rich_text"))
        doc.add_paragraph(text, style='List Number')
    elif block_type == "to_do":
        text = extract_rich_text(data.get("rich_text"))
        checked = data.get("checked", False)
        mark = "[x]" if checked else "[ ]"
        add_text_paragraph(f"{mark} {text}")
    elif block_type == "quote":
        text = extract_rich_text(data.get("rich_text"))
        add_text_paragraph(f"“{text}”", style='Quote')
    elif block_type == "code":
        text = extract_rich_text(data.get("rich_text"))
        # language = data.get("language", "")
        # Very simple code block representation
        p = doc.add_paragraph(style='No Spacing')
        runner = p.add_run(text)
        runner.font.name = 'Courier New'
    elif block_type == "image":
        image_data = data.get(data.get("type", ""), {})
        image_url = image_data.get("url", "")
        if image_url:
            logging.info(f"Processing image block {block['id']} with URL: {image_url[:50]}...")
            stream = download_image_to_stream(image_url, block_id=block['id'])
            
            if stream:
                try:
                    doc.add_picture(stream, width=Inches(5.0))
                except Exception as e:
                    logging.warning(f"Could not add picture to docx: {e}")
    else:
        text = extract_rich_text(data.get("rich_text"))
        if text:
             add_text_paragraph(text)

    if block.get("has_children"):
        children = list_block_children(notion, block["id"])
        child_depth = depth + 1
        for child in children:
            write_block_to_docx(doc, child, notion, depth=child_depth)


def page_to_docx(notion: Client, page: Dict[str, Any]) -> Document:
    doc = Document()
    doc.add_heading(get_page_title(page), 0) # Title
    
    blocks = list_block_children(notion, page["id"])
    for block in blocks:
        write_block_to_docx(doc, block, notion)
        
    return doc


def page_to_markdown(notion: Client, page: Dict[str, Any]) -> str:
    blocks = list_block_children(notion, page["id"])
    lines = []
    for block in blocks:
        lines.extend(block_to_markdown(notion, block))
    return "\n".join(lines).strip() + "\n"


def get_page_title(page: Dict[str, Any]) -> str:
    for prop in page.get("properties", {}).values():
        if prop.get("type") == "title":
            return extract_rich_text(prop.get("title"))
    return "Untitled"


def get_drive_service():
    creds = None
    token_refresh_failed = False  # Track if we had a token that failed to refresh

    # 1. Try loading from Environment Variable (OAuth token as JSON string - for GitHub Actions)
    env_token = os.getenv("GOOGLE_TOKEN_JSON")
    if env_token:
        try:
            logging.info("Loading Google Creds from GOOGLE_TOKEN_JSON env var...")
            info = json.loads(env_token)
            creds = Credentials.from_authorized_user_info(info, DRIVE_SCOPES)
        except Exception as e:
            logging.warning(f"Failed to load token from env: {e}")

    # 2. Try loading from local token.json file (Local Development)
    if not creds and os.path.exists('token.json'):
        try:
            logging.info("Loading Google Creds from local token.json...")
            creds = Credentials.from_authorized_user_file('token.json', DRIVE_SCOPES)
        except Exception as e:
            logging.warning(f"Failed to load token.json: {e}")

    # 3. Refresh if expired
    if creds and creds.expired and creds.refresh_token:
        try:
            logging.info("Refreshing expired OAuth token...")
            creds.refresh(Request())
        except Exception as e:
            logging.warning(f"Token refresh failed: {e}")
            token_refresh_failed = True
            creds = None

    # 4. Fallback to Service Account (Only for Workspace/Enterprise users with quota)
    if not creds:
        sa_path = os.getenv("GOOGLE_APPLICATION_CREDENTIALS")
        if sa_path and os.path.exists(sa_path):
            try:
                logging.info(f"Loading Google Creds from service account file: {sa_path}")
                creds = service_account.Credentials.from_service_account_file(
                    sa_path, scopes=DRIVE_SCOPES)
            except Exception as e:
                logging.warning(f"Failed to load service account: {e}")

    # 5. Interactive Login (Last Resort for Local)
    if not creds:
        if os.getenv("GITHUB_ACTIONS"):
            # Give clear error message for GitHub Actions
            if token_refresh_failed:
                raise RuntimeError(
                    "OAuth token refresh failed - your refresh token has expired or been revoked.\n\n"
                    "To fix this:\n"
                    "1. Run 'python setup_oauth.py' locally to generate a new token\n"
                    "2. Copy the contents of token.json\n"
                    "3. Update GOOGLE_TOKEN_JSON secret in GitHub repo settings\n\n"
                    "Note: If your OAuth app is in 'Testing' mode, tokens expire after 7 days.\n"
                    "Consider publishing your app or adding your email to the test users list."
                )
            else:
                raise RuntimeError(
                    "No valid Google credentials found for GitHub Actions.\n\n"
                    "Please set the GOOGLE_TOKEN_JSON secret with valid OAuth token JSON."
                )

        if not os.path.exists('client_secret.json'):
            logging.warning("No credentials found. Please run 'setup_oauth.py' to generate token.json.")
            return None

        logging.info("No token found. Initiating interactive login...")
        flow = InstalledAppFlow.from_client_secrets_file(
            'client_secret.json', DRIVE_SCOPES)
        creds = flow.run_local_server(port=0)

        # Save the new token locally
        with open('token.json', 'w') as token:
            token.write(creds.to_json())

    return build('drive', 'v3', credentials=creds)


def escape_drive_query_value(value: str) -> str:
    return value.replace("'", "\\'")


def find_drive_file(service, folder_id: str, filename: str) -> Optional[Dict[str, Any]]:
    safe_name = escape_drive_query_value(filename)
    safe_folder = escape_drive_query_value(folder_id)
    query = f"name = '{safe_name}' and '{safe_folder}' in parents and trashed = false"
    response = (
        service.files()
        .list(q=query, fields="files(id, name)", pageSize=1)
        .execute()
    )
    files = response.get("files", [])
    return files[0] if files else None


def upload_to_drive(service, folder_id: str, filename: str, local_path: Path) -> str:
    existing = find_drive_file(service, folder_id, filename)
    # Correct MIME type for .docx
    media = MediaFileUpload(str(local_path), mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document", resumable=True)

    if existing and not OVERWRITE_EXISTING:
        return "skipped"

    if existing and OVERWRITE_EXISTING:
        service.files().update(fileId=existing["id"], media_body=media).execute()
        return "updated"

    file_metadata = {"name": filename, "parents": [folder_id]}
    service.files().create(body=file_metadata, media_body=media, fields="id").execute()
    return "created"


def ensure_required_config() -> None:
    missing = []
    if not NOTION_TOKEN:
        missing.append("NOTION_TOKEN")
    if not DATABASE_ID:
        missing.append("DATABASE_ID")
    if not DRIVE_FOLDER_ID:
        missing.append("DRIVE_FOLDER_ID")
    # Credential check: Need at least one valid method
    has_creds = False
    
    # 1. OAuth Token from Env
    if os.getenv("GOOGLE_TOKEN_JSON"):
        has_creds = True
    # 2. Local OAuth Token
    elif Path(BASE_DIR / "token.json").exists():
        has_creds = True
    # 3. Service Account File
    elif SERVICE_ACCOUNT_FILE and Path(SERVICE_ACCOUNT_FILE).exists():
        has_creds = True
        
    if not has_creds:
        missing.append("GOOGLE_CREDENTIALS (GOOGLE_TOKEN_JSON env or token.json or valid GOOGLE_APPLICATION_CREDENTIALS)")

    if missing:
        raise RuntimeError(f"Missing required config: {', '.join(missing)}")


def main() -> None:
    ensure_required_config()
    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)

    notion = Client(auth=NOTION_TOKEN)
    drive_service = get_drive_service()

    last_sync_time = load_last_sync_time()
    pages = query_all_pages(notion, DATABASE_ID, last_sync_time)

    if not pages:
        logging.info("No pages to sync.")
        return

    failures = 0

    for page in pages:
        page_title = get_page_title(page)
        last_edited = page.get("last_edited_time", "")
        date_stamp = (last_edited or now_utc_iso())[:10]
        # Changed extension to .docx
        filename = sanitize_filename(f"{page_title}_{date_stamp}.docx")
        local_path = EXPORTS_DIR / filename

        try:
            # Generate DOCX instead of Markdown
            doc = page_to_docx(notion, page)
            doc.save(local_path)
            
            result = upload_to_drive(drive_service, DRIVE_FOLDER_ID, filename, local_path)
            logging.info(f"Synced ({result}): {page_title}")
        except HttpError as exc:
            failures += 1
            if exc.resp.status == 403:
                logging.error(
                    "Failed: Ensure Service Account email is added as Editor to the Folder."
                )
            else:
                logging.error(f"Failed: {page_title} ({exc})")
        except Exception as exc:
            failures += 1
            logging.error(f"Failed: {page_title} ({exc})")

    if failures == 0:
        save_last_sync_time(now_utc_iso())
    else:
        logging.warning("Sync completed with failures; last_sync_time not updated.")


if __name__ == "__main__":
    main()
